from pathlib import Path
from pypdf import PdfReader
from pptx import Presentation
import pandas as pd

BASE_DIR = Path(__file__).resolve().parent.parent.parent

PDF_DIR = BASE_DIR / "data" / "documents" / "raw" / "pdf"
PPT_DIR = BASE_DIR / "data" / "documents" / "raw" / "ppt"
VIDEO_DIR = BASE_DIR / "data" / "documents" / "raw" / "video"

OUTPUT_DIR = BASE_DIR / "outputs" / "documents"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

def extract_pdf_text(pdf_path):
    text = ""
    try:
        reader = PdfReader(pdf_path)
        for page in reader.pages[:3]:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    except Exception as e:
        text = f"ERROR: {e}"
    return text.strip()

def extract_ppt_text(ppt_path):
    text = ""
    try:
        prs = Presentation(ppt_path)
        for slide in prs.slides[:5]:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text += shape.text + "\n"
    except Exception as e:
        text = f"ERROR: {e}"
    return text.strip()

def main():
    rows = []

    for pdf_file in PDF_DIR.glob("*.pdf"):
        rows.append({
            "original_file_name": pdf_file.name,
            "file_type": "PDF",
            "text_preview": extract_pdf_text(pdf_file)[:3000]
        })

    for ppt_file in PPT_DIR.glob("*.pptx"):
        rows.append({
            "original_file_name": ppt_file.name,
            "file_type": "PPTX",
            "text_preview": extract_ppt_text(ppt_file)[:3000]
        })

    for video_file in VIDEO_DIR.glob("*"):
        rows.append({
            "original_file_name": video_file.name,
            "file_type": "VIDEO",
            "text_preview": "video file"
        })

    df = pd.DataFrame(rows)
    df.to_excel(OUTPUT_DIR / "document_text_extraction.xlsx", index=False)

    print("\n=== Extracted Documents Preview ===")
    print(df.to_string(index=False))
    print("\ndocument_text_extraction.xlsx created")

if __name__ == "__main__":
    main()