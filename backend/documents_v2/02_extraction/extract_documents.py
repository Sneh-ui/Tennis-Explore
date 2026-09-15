import json
import sys
import zipfile
from datetime import datetime
from pathlib import Path

from pypdf import PdfReader
from pptx import Presentation


# ---------------------------------------------------------
# PROJECT CONFIGURATION
# ---------------------------------------------------------

CURRENT_FILE = Path(__file__).resolve()
DOCUMENTS_V2_DIR = CURRENT_FILE.parent.parent

if str(DOCUMENTS_V2_DIR) not in sys.path:
    sys.path.insert(0, str(DOCUMENTS_V2_DIR))

from config import (
    PDF_TEST_DIR,
    PPTX_TEST_DIR,
    EXTRACTION_OUTPUT_DIR,
)


# ---------------------------------------------------------
# HELPERS
# ---------------------------------------------------------

def clean_text(text):
    """
    Perform basic text cleanup without changing
    the meaning of the extracted document.
    """
    if not text:
        return ""

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    return "\n".join(lines)


def save_json(data, output_path):
    """Save extracted document data as UTF-8 JSON."""

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with output_path.open(
        "w",
        encoding="utf-8",
    ) as file:
        json.dump(
            data,
            file,
            ensure_ascii=False,
            indent=2,
        )


# ---------------------------------------------------------
# PDF EXTRACTION
# ---------------------------------------------------------

def extract_pdf(file_path):
    """
    Extract text from EVERY page of a PDF.

    Unlike the Semester 1 prototype, there is
    deliberately no five-page restriction.
    """

    print()
    print(f"Extracting PDF: {file_path.name}")

    document = {
        "source_file": file_path.name,
        "source_path": str(file_path),
        "file_type": "pdf",
        "file_size_bytes": file_path.stat().st_size,
        "total_pages": 0,
        "pages_with_text": 0,
        "empty_pages": 0,
        "extraction_errors": [],
        "pages": [],
    }

    try:
        reader = PdfReader(str(file_path))

        document["total_pages"] = len(reader.pages)

        for page_number, page in enumerate(
            reader.pages,
            start=1,
        ):
            try:
                raw_text = page.extract_text() or ""
                text = clean_text(raw_text)

                has_text = bool(text)

                if has_text:
                    document["pages_with_text"] += 1
                else:
                    document["empty_pages"] += 1

                document["pages"].append(
                    {
                        "page_number": page_number,
                        "has_text": has_text,
                        "character_count": len(text),
                        "word_count": len(text.split()),
                        "text": text,
                    }
                )

            except Exception as exc:
                document["extraction_errors"].append(
                    {
                        "page_number": page_number,
                        "error_type": type(exc).__name__,
                        "message": str(exc),
                    }
                )

                document["pages"].append(
                    {
                        "page_number": page_number,
                        "has_text": False,
                        "character_count": 0,
                        "word_count": 0,
                        "text": "",
                    }
                )

        # Detect PDFs that opened correctly but contain no extractable text.
        # These are likely scanned/image-based PDFs and may require OCR.

        if (
            document["total_pages"] > 0
            and document["pages_with_text"] == 0
        ):
            document["requires_ocr"] = True
            document["status"] = "needs_ocr"

        else:
            document["requires_ocr"] = False

            document["status"] = (
                "success"
                if not document["extraction_errors"]
                else "completed_with_errors"
            )

    except Exception as exc:
        document["status"] = "failed"

        document["extraction_errors"].append(
            {
                "page_number": None,
                "error_type": type(exc).__name__,
                "message": str(exc),
            }
        )

    return document


# ---------------------------------------------------------
# POWERPOINT EXTRACTION
# ---------------------------------------------------------

def extract_text_from_shape(shape):
    """
    Extract text from a PowerPoint shape.

    Supports:
    - normal text boxes/placeholders
    - tables
    - grouped/nested shapes
    """

    text_parts = []

    # -----------------------------------------------------
    # 1. Normal text boxes / placeholders
    # -----------------------------------------------------
    if hasattr(shape, "has_text_frame") and shape.has_text_frame:
        shape_text = clean_text(shape.text)

        if shape_text:
            text_parts.append(shape_text)

    # -----------------------------------------------------
    # 2. Tables
    # -----------------------------------------------------
    if hasattr(shape, "has_table") and shape.has_table:

        for row in shape.table.rows:
            row_parts = []

            for cell in row.cells:
                cell_text = clean_text(cell.text)

                if cell_text:
                    row_parts.append(cell_text)

            if row_parts:
                text_parts.append(" | ".join(row_parts))

    # -----------------------------------------------------
    # 3. Grouped / nested shapes
    # -----------------------------------------------------
    if hasattr(shape, "shapes"):

        for nested_shape in shape.shapes:
            nested_text = extract_text_from_shape(
                nested_shape
            )

            if nested_text:
                text_parts.append(nested_text)

    return "\n".join(text_parts)


def extract_pptx(file_path):
    """
    Extract text from every slide of a PPTX while
    preserving slide numbers.

    Includes text boxes, tables, and grouped shapes.
    """

    print()
    print(f"Extracting PPTX: {file_path.name}")

    document = {
        "source_file": file_path.name,
        "source_path": str(file_path),
        "file_type": "pptx",
        "file_size_bytes": file_path.stat().st_size,
        "total_slides": 0,
        "slides_with_text": 0,
        "empty_slides": 0,
        "extraction_errors": [],
        "slides": [],
    }

    if not zipfile.is_zipfile(file_path):
        document["status"] = "invalid_pptx"

        document["extraction_errors"].append(
            {
                "slide_number": None,
                "error_type": "InvalidPptxPackage",
                "message": (
                    "File is not a valid PPTX/ZIP package."
                ),
            }
        )

        return document

    try:
        presentation = Presentation(str(file_path))

        document["total_slides"] = len(
            presentation.slides
        )

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            try:
                text_parts = []

                for shape in slide.shapes:

                    shape_text = extract_text_from_shape(
                        shape
                    )

                    if shape_text:
                        text_parts.append(shape_text)

                text = clean_text(
                    "\n".join(text_parts)
                )

                has_text = bool(text)

                if has_text:
                    document["slides_with_text"] += 1
                else:
                    document["empty_slides"] += 1

                document["slides"].append(
                    {
                        "slide_number": slide_number,
                        "has_text": has_text,
                        "character_count": len(text),
                        "word_count": len(text.split()),
                        "text": text,
                    }
                )

            except Exception as exc:
                document["extraction_errors"].append(
                    {
                        "slide_number": slide_number,
                        "error_type": type(exc).__name__,
                        "message": str(exc),
                    }
                )

                document["slides"].append(
                    {
                        "slide_number": slide_number,
                        "has_text": False,
                        "character_count": 0,
                        "word_count": 0,
                        "text": "",
                    }
                )

        document["status"] = (
            "success"
            if not document["extraction_errors"]
            else "completed_with_errors"
        )

    except Exception as exc:
        document["status"] = "failed"

        document["extraction_errors"].append(
            {
                "slide_number": None,
                "error_type": type(exc).__name__,
                "message": str(exc),
            }
        )

    return document


# ---------------------------------------------------------
# PROCESS FILE
# ---------------------------------------------------------

def process_document(
    file_path,
    output_stem=None,
    source_file_name=None,
    source_metadata=None,
):
    """
    Extract one supported document.

    Optional production arguments:
    - output_stem:
        unique name used for the output JSON
    - source_file_name:
        original source filename to preserve in metadata
    - source_metadata:
        additional source metadata such as S3 key,
        ETag, last-modified timestamp, etc.

    Existing local/test usage continues to work
    without providing any optional arguments.
    """

    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        result = extract_pdf(file_path)

    elif suffix == ".pptx":
        result = extract_pptx(file_path)

    else:
        print(
            f"Skipping unsupported file: "
            f"{file_path.name}"
        )
        return None

    # Preserve the true/original source filename
    # when production ingestion temporarily renames
    # the downloaded file.
    if source_file_name:
        result["source_file"] = (
            source_file_name
        )

    # Add production/source metadata without
    # affecting existing extraction fields.
    if source_metadata:
        result["source_metadata"] = (
            source_metadata
        )

    result["extracted_at"] = (
        datetime.now()
        .astimezone()
        .isoformat()
    )

    # Existing behaviour when output_stem is absent.
    if output_stem is None:
        output_stem = file_path.stem

    result["output_stem"] = (
        output_stem
    )

    output_name = (
        f"{output_stem}"
        "_extracted.json"
    )

    output_path = (
        EXTRACTION_OUTPUT_DIR
        / output_name
    )

    save_json(
        result,
        output_path,
    )

    print(
        f"Status: {result['status']}"
    )

    print(
        f"Saved: {output_path}"
    )

    return result


# ---------------------------------------------------------
# SUMMARY
# ---------------------------------------------------------

def print_summary(results):
    print()
    print("=" * 70)
    print("EXTRACTION SUMMARY")
    print("=" * 70)

    successful = 0
    needs_ocr = 0
    with_errors = 0
    failed = 0

    for result in results:

        status = result["status"]

        if status == "success":
            successful += 1

        elif status == "needs_ocr":
            needs_ocr += 1

        elif status == "completed_with_errors":
            with_errors += 1

        elif status == "failed":
            failed += 1

        print()
        print(
            f"File   : {result['source_file']}"
        )
        print(
            f"Type   : {result['file_type'].upper()}"
        )
        print(
            f"Status : {status}"
        )

        if result["file_type"] == "pdf":
            print(
                f"Pages  : "
                f"{result['total_pages']}"
            )
            print(
                f"Text   : "
                f"{result['pages_with_text']} "
                "pages"
            )
            print(
                f"Empty  : "
                f"{result['empty_pages']} "
                "pages"
            )

        else:
            print(
                f"Slides : "
                f"{result['total_slides']}"
            )
            print(
                f"Text   : "
                f"{result['slides_with_text']} "
                "slides"
            )
            print(
                f"Empty  : "
                f"{result['empty_slides']} "
                "slides"
            )

        print(
            f"Errors : "
            f"{len(result['extraction_errors'])}"
        )

    print()
    print("-" * 70)
    print(f"Successful            : {successful}")
    print(f"Needs OCR             : {needs_ocr}")
    print(f"Completed with errors : {with_errors}")
    print(f"Failed                : {failed}")
    print(f"Total processed       : {len(results)}")
    print("-" * 70)


# ---------------------------------------------------------
# MAIN
# ---------------------------------------------------------

def main():

    print()
    print("=" * 70)
    print("TENNIS EXPLORE - V2 DOCUMENT EXTRACTION")
    print("=" * 70)

    results = []

    # PDFs
    if PDF_TEST_DIR.exists():

        pdf_files = sorted(
            PDF_TEST_DIR.glob("*.pdf")
        )

        print(
            f"\nPDF files found: "
            f"{len(pdf_files)}"
        )

        for file_path in pdf_files:
            result = process_document(file_path)

            if result:
                results.append(result)

    else:
        print(
            f"\nPDF directory not found: "
            f"{PDF_TEST_DIR}"
        )

    # PowerPoints
    if PPTX_TEST_DIR.exists():

        pptx_files = sorted(
            PPTX_TEST_DIR.glob("*.pptx")
        )

        print(
            f"\nPPTX files found: "
            f"{len(pptx_files)}"
        )

        for file_path in pptx_files:
            result = process_document(file_path)

            if result:
                results.append(result)

    else:
        print(
            f"\nPPTX directory not found: "
            f"{PPTX_TEST_DIR}"
        )

    if not results:
        print()
        print("No supported documents were processed.")
        return

    print_summary(results)


if __name__ == "__main__":
    main()